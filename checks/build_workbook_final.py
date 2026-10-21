"""Rebuild the paper's data workbook (meta-analysis-data.xlsx) after the author check: estimates sheet from
results/tables/estimates_with_effect_sizes.csv; panels_and_channels without removed studies, n_obs from the corrected data,
kept_one_per_panel = largest N in the panel (ties keep the first-pass choice). First-pass workbook: checks/firstpass/."""
import sys, csv, math, re
import openpyxl
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
R = r"D:/Research-Publishing/14-fintech-emissions-meta/"
wb = openpyxl.load_workbook(R + "checks/firstpass/meta-analysis-data.xlsx")
ws = wb["estimates"]; hdr = [c.value for c in ws[1]]
num_int = {'primary','in_main_table','year','venue_flag','sample_mid_year','n_obs','n_units','k_reg','k_reg_reported','page','inconsistent_flag','n_inferred_flag','primary_fallback_flag','retracted_flag','sample_overlap_flag'}
num_float = {'b','se','t','p'}
rows = list(csv.DictReader(open(R + "results/tables/estimates_with_effect_sizes.csv", encoding="utf-8")))
ws.delete_rows(2, ws.max_row - 1)
def conv(h, v):
    if v in ("", "NA", None): return None
    if h in num_int:
        try: return int(float(v))
        except ValueError: return v
    if h in num_float:
        try: return float(v)
        except ValueError: return v
    return ILLEGAL_CHARACTERS_RE.sub(" ", v)
for r in rows:
    ws.append([round(float(r["ri"]), 6) if h == "r" else round(float(r["zi"]), 6) if h == "z" else round(math.sqrt(float(r["vi"])), 6) if h == "se_z" else conv(h, r.get(h, "")) for h in hdr])
ps = wb["panels_and_channels"]; ph = [c.value for c in ps[1]]; I = {h: i for i, h in enumerate(ph)}
prim = {r["study_id"]: r for r in rows if r["primary"] == "1"}
for rr in sorted([row[0].row for row in ps.iter_rows(min_row=2) if row[0].value not in prim], reverse=True): ps.delete_rows(rr)
for row in ps.iter_rows(min_row=2):
    row[I["n_obs"]].value = int(float(prim[row[0].value]["n_obs"]))
data = [[c.value for c in r] for r in ps.iter_rows(min_row=2)]
for rule in "AB":
    col = I[f"panel_rule_{rule}"]; kc = I[f"kept_one_per_panel_{rule}"]; groups = {}
    for j, d in enumerate(data): groups.setdefault(d[0] if d[col] == "own panel" else d[col], []).append(j)
    for p, js in groups.items():
        mx = max(float(data[j][I["n_obs"]]) for j in js); top = [j for j in js if float(data[j][I["n_obs"]]) == mx]
        prev = [j for j in top if int(data[j][kc]) == 1]
        best = prev[0] if prev else top[0]
        for j in js: ps.cell(row=j + 2, column=kc + 1).value = 1 if j == best else 0
cb = wb["columns"]
for r in cb.iter_rows(min_row=2):
    if r[1].value:
        r[1].value = r[1].value.replace("all 55 included studies", "all 54 included studies").replace("(91 main-table estimates from 55 studies)", "(89 main-table estimates from 54 studies, after the author check of 22 Oct 2026)")
wb.save(R + "data/workbook/meta-analysis-data.xlsx")
print("estimates", ws.max_row - 1, "panels", ps.max_row - 1)
